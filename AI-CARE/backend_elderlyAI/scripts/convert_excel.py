import openpyxl
import json
import sys
from pathlib import Path

def main():
    # Relative path calculation
    script_dir = Path(__file__).resolve().parent
    base_dir = script_dir.parent.parent  # AI-CARE root

    excel_path = base_dir / "AI_CARE_Database.xlsx"
    out_json_path = base_dir / "frontend_elderlyAI" / "src" / "data" / "patientsFromExcel.json"

    if not excel_path.exists():
        print(f"Error: Could not find Excel database at {excel_path}")
        sys.exit(1)

    wb = openpyxl.load_workbook(str(excel_path), data_only=True)

    # 1. Read Patients
    patients_ws = wb['Patients']
    p_rows = list(patients_ws.iter_rows(values_only=True))
    p_headers = [str(h) for h in p_rows[0]]
    patients_dict = {}

    for row in p_rows[1:]:
        if not row or not row[0]:
            continue
        r = dict(zip(p_headers, row))
        p_id = str(r['patient_id']).strip()
        phone_raw = str(r.get('phone', '')) if r.get('phone') is not None else ''
        if phone_raw.endswith('.0'):
            phone_raw = phone_raw[:-2]
        phone_formatted = f'0{phone_raw}' if not phone_raw.startswith('0') else phone_raw

        patients_dict[p_id] = {
            'patient_id': p_id,
            'id': p_id,
            'patient_code': p_id,
            'device_id': str(r.get('device_id', '')),
            'full_name': str(r.get('name', '')),
            'fullName': str(r.get('name', '')),
            'name': str(r.get('name', '')),
            'username': str(r.get('username', '')),
            'email': str(r.get('email', '')),
            'role': str(r.get('role', 'Patient')),
            'status': str(r.get('status', 'Hoạt động')),
            'age': int(r.get('age', 70)) if r.get('age') is not None else 70,
            'gender': str(r.get('gender', 'Nam')),
            'phone': phone_formatted,
            'height_cm': int(r.get('height_cm', 160)) if r.get('height_cm') is not None else 160,
            'weight_kg': int(r.get('weight_kg', 60)) if r.get('weight_kg') is not None else 60,
            'blood_group': str(r.get('blood_group', 'O+')),
            'allergy': str(r.get('allergy', 'Không')),
            'address': f"Khu vực giám sát thiết bị {r.get('device_id', 'AI')}",
            'medical_history': 'Theo dõi định kỳ lão khoa',
            'medicalConditions': 'Theo dõi định kỳ lão khoa',
            'caregiver_name': 'Chưa cập nhật',
            'caregiver_relation': 'Người thân',
            'caregiver_age': 42,
            'caregiver_phone': '0987654321',
            'caregiver_email': 'family@elderly.ai',
            'relativeName': 'Chưa cập nhật',
            'relativeRelation': 'Người thân',
            'relativeAge': 42,
            'relativePhone': '0987654321',
            'relativeEmail': 'family@elderly.ai',
            'doctor_name': 'BS. Nguyễn Thanh Tùng'
        }

    # 2. Read Caregivers
    if 'Caregivers' in wb.sheetnames:
        cg_ws = wb['Caregivers']
        cg_rows = list(cg_ws.iter_rows(values_only=True))
        cg_headers = [str(h) for h in cg_rows[0]]
        for row in cg_rows[1:]:
            if not row or not row[0]: continue
            r = dict(zip(cg_headers, row))
            p_id = str(r['patient_id']).strip()
            if p_id in patients_dict:
                cg_phone_raw = str(r.get('caregiver_phone', '')) if r.get('caregiver_phone') is not None else ''
                if cg_phone_raw.endswith('.0'): cg_phone_raw = cg_phone_raw[:-2]
                cg_phone = f'0{cg_phone_raw}' if not cg_phone_raw.startswith('0') else cg_phone_raw
                cg_name = str(r.get('caregiver_name', 'Người thân'))
                cg_relation = str(r.get('caregiver_relation', 'Người thân'))
                cg_email = str(r.get('caregiver_email', 'family@elderly.ai'))

                patients_dict[p_id]['caregiver_name'] = cg_name
                patients_dict[p_id]['relativeName'] = cg_name
                patients_dict[p_id]['caregiver_relation'] = cg_relation
                patients_dict[p_id]['relativeRelation'] = cg_relation
                patients_dict[p_id]['caregiver_phone'] = cg_phone
                patients_dict[p_id]['relativePhone'] = cg_phone
                patients_dict[p_id]['caregiver_email'] = cg_email
                patients_dict[p_id]['relativeEmail'] = cg_email

    # 3. Read Doctors
    if 'Doctors' in wb.sheetnames:
        doc_ws = wb['Doctors']
        doc_rows = list(doc_ws.iter_rows(values_only=True))
        doc_headers = [str(h) for h in doc_rows[0]]
        for row in doc_rows[1:]:
            if not row or not row[0]: continue
            r = dict(zip(doc_headers, row))
            p_id = str(r['patient_id']).strip()
            if p_id in patients_dict:
                doc_name = str(r.get('doctor_name', 'BS. Nguyễn Thanh Tùng'))
                patients_dict[p_id]['doctor_name'] = doc_name

    # 4. Read Health_Records
    if 'Health_Records' in wb.sheetnames:
        hr_ws = wb['Health_Records']
        hr_rows = list(hr_ws.iter_rows(values_only=True))
        hr_headers = [str(h) for h in hr_rows[0]]
        for row in hr_rows[1:]:
            if not row or not row[0]: continue
            r = dict(zip(hr_headers, row))
            p_id = str(r['patient_id']).strip()
            if p_id in patients_dict:
                disease = str(r.get('disease', 'Theo dõi định kỳ'))
                patients_dict[p_id]['medical_history'] = disease
                patients_dict[p_id]['medicalConditions'] = disease

    patient_list = list(patients_dict.values())
    out_json_path.parent.mkdir(parents=True, exist_ok=True)
    with open(out_json_path, 'w', encoding='utf-8') as f:
        json.dump(patient_list, f, ensure_ascii=False, indent=2)

    print(f"Successfully exported {len(patient_list)} patients to {out_json_path}")

if __name__ == '__main__':
    main()
